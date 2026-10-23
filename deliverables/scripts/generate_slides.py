import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette: Deep Navy, Electric Blue, Slate, White, Coral/Amber
    BG_COLOR = RGBColor(10, 14, 23)        # #0a0e17
    CARD_BG = RGBColor(26, 34, 52)         # #1a2234
    ACCENT_BLUE = RGBColor(59, 130, 246)   # #3b82f6
    TEXT_WHITE = RGBColor(248, 250, 252)   # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94a3b8
    ACCENT_AMBER = RGBColor(245, 158, 11)  # #f59e0b
    ACCENT_GREEN = RGBColor(16, 185, 129)  # #10b981
    BORDER_COLOR = RGBColor(46, 61, 91)    # #2e3d5b

    blank_layout = prs.slide_layouts[6]

    def add_base_slide(title_text, subtitle_text, slide_num):
        slide = prs.slides.add_slide(blank_layout)
        
        # Dark Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR

        # Header bar
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "Arial"

        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(13)
        p2.font.color.rgb = ACCENT_BLUE
        p2.font.name = "Arial"

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        ftf = footer_box.text_frame
        ftf.word_wrap = True
        fp = ftf.paragraphs[0]
        fp.text = f"SpendNest Revenue Intelligence Inbox • Submission Slide {slide_num} of 5"
        fp.font.size = Pt(10)
        fp.font.color.rgb = TEXT_MUTED

        return slide

    def add_card(slide, left, top, width, height, title, body_bullets, accent_color=ACCENT_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        tx_box = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = accent_color
        p.font.name = "Arial"

        for b in body_bullets:
            p_b = tf.add_paragraph()
            p_b.text = "• " + b
            p_b.font.size = Pt(11)
            p_b.font.color.rgb = TEXT_WHITE
            p_b.font.name = "Arial"
            p_b.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: Manager Problem & 3 Core Findings
    # -------------------------------------------------------------
    s1 = add_base_slide(
        "Manager Problem Solved & Three Critical Findings",
        "Taming noisy multi-channel signals into actionable pipeline intelligence with traceable citations",
        1
    )
    add_card(s1, 0.8, 1.8, 3.7, 4.8, "The Manager Problem", [
        "A sales manager receives fragmented calls, meetings, and AI notes with silent audio and delayed revisions.",
        "Blindly relying on stored metadata leads to false pipeline numbers and misdirected coaching.",
        "Our Solution: Connects metadata directly to turn-level transcript evidence, separates call reps from calendar hosts, and preserves manual routing corrections across updates."
    ], ACCENT_BLUE)

    add_card(s1, 4.8, 1.8, 7.7, 4.8, "Three Crucial Source Findings", [
        "Finding 1 (M02 - LedgerBridge, Turn 1-3): Stored meeting classified as sales opportunity, but SpendNest is actually evaluating LedgerBridge as an e-invoicing supplier. Product choice: Routed to 'Vendor Note', preventing false ARR inflation.",
        "Finding 2 (M07 - Willowfield Trust, Turn 2-8): Late transcript in revision 2 revealed severe production UI bug (silent failure of payment description save + recreate attachment block). Product choice: Routed immediately to P1 Engineering Bug ticket with turn citations.",
        "Finding 3 (C09 - Vedant, Exact Quote): Prospect Vedant referred rep to business head Rishabh with strict mandate: 'No, please don\\'t use my name.' Product choice: Enforced explicit compliance coaching alert preventing rep outreach violation."
    ], ACCENT_AMBER)

    # -------------------------------------------------------------
    # SLIDE 2: Interface Choices & Deliberate Omissions
    # -------------------------------------------------------------
    s2 = add_base_slide(
        "Interface Choices & Deliberate Omissions",
        "Designed for rapid executive decision-making: high signal density, clear boundaries, zero vanity fluff",
        2
    )
    add_card(s2, 0.8, 1.8, 5.7, 4.8, "What We Intentionally Built", [
        "Interactive Replay Control Bar: 1-click execution of the entire replay contract (initial import, shuffle test, override, update upsert, regression test).",
        "Role Distinction: Strictly isolates Call Outbound Reps (Arun, Meera, Leela, Kiran, Dev) from Meeting Calendar Hosts (team08, team09, etc.).",
        "Three-Pane Evidence Workspace: Raw envelope JSON, verbatim transcript turns with timestamps, and candidate-derived analysis side-by-side.",
        "Live Manifest & Export Engine: One-click export of machine-readable JSON/CSV manifests and organized markdown route files."
    ], ACCENT_GREEN)

    add_card(s2, 6.8, 1.8, 5.7, 4.8, "What We Deliberately Left Out", [
        "No Fictitious Pipeline ARR Totals: The data lacks verified contract values and win probabilities. We refused to invent arbitrary deal sizes.",
        "No Automated Overwrite of Human Decisions: When revision 2 arrived for M07, the app preserved the manager's manual route rather than wiping it.",
        "No Substantive AI Analysis for Silent Intake: 26 records had null transcripts (B01-B20, M05, etc.). We logged them as intake, refusing to fabricate summaries.",
        "No Multi-Date Filters: All 44 records occurred on 2026-08-21. A date picker adds clutter; we prioritized source-type, rep, and route filters."
    ], ACCENT_AMBER)

    # -------------------------------------------------------------
    # SLIDE 3: Routing Taxonomy & Difficult Routing Case
    # -------------------------------------------------------------
    s3 = add_base_slide(
        "Routing Architecture & Difficult Routing Case",
        "Differentiated destinations matching operational reality, highlighted by M07's late arrival",
        3
    )
    add_card(s3, 0.8, 1.8, 5.7, 4.8, "Six-Pillar Routing Taxonomy", [
        "1. Sales Coaching & Guidance (6): Evaluates objection handling (C01), competitive intel (C05, C06), and compliance (C09).",
        "2. Deal Next Steps (4): Actionable commercial follow-ups with stated owners (C03 Milan, C04 Lanternpath, C08 Raghav, M03 Asterloom).",
        "3. Customer Support & Escalations (3): Critical operational issues (M04 Card Migration Sunset, M07 P1 Bug, M09 Onboarding).",
        "4. Partner & Channel Onboarding (2): Referral agreements and commissions (C07 Willowbank, M01 Harborpath).",
        "5. Vendor / Internal Ops Notes (2): Procurement & training (M02 LedgerBridge supplier, M10 billing training).",
        "6. Voicemail / Gatekeeper & Unusable (27): Clean separation of non-conversational dials and administrative transfers (C02, B01-B20)."
    ], ACCENT_BLUE)

    add_card(s3, 6.8, 1.8, 5.7, 4.8, "The Difficult Case: M07 (Willowfield)", [
        "Initial Delivery: M07 arrived with summary: null and transcript_payload: null (Rev 1). Appeared as silent intake.",
        "Manager Action: Manager proactively overrides route based on customer communication to 'Customer Support / Critical Bug Escalation'.",
        "Update Delivery: Batch 02_update arrives with Revision 2 and 27 turns of screen-share troubleshooting.",
        "The Resolution: The system preserves the manager's manual route, flags the analysis as 'Revision Updated (Pending Review)', highlights turns 2-8 evidence, and avoids creating a duplicate conversation record."
    ], ACCENT_AMBER)

    # -------------------------------------------------------------
    # SLIDE 4: Verification of Numbers & Analytical Conclusions
    # -------------------------------------------------------------
    s4 = add_base_slide(
        "Verification of Key Numbers & Analytical Evidence",
        "Mathematical rigor and verifiable citations: testing denominators, revisions, and claims",
        4
    )
    add_card(s4, 0.8, 1.8, 5.7, 4.8, "Verifying Headline Numbers", [
        "Total Ingested Records (44): Verified 30 calls + 14 meetings. Ingesting 02_update (1 record) does NOT increase count to 45 because (source_type, source_id) matches M07.",
        "Transcript Coverage (17 -> 18): Initially 10 calls + 7 meetings = 17 (38.6%). Post-update: M07 adds 1 meeting = 18 (40.9%). Denominator is strictly 44 logical records.",
        "Idempotency Test: Shuffling initial batch and re-importing yields 0 duplicate records and 0 inflated transcripts.",
        "Revision Precedence: Re-importing 01_initial after 02_update safely ignores Revision 1 for M07; Revision 2 and 27 turns remain intact."
    ], ACCENT_GREEN)

    add_card(s4, 6.8, 1.8, 5.7, 4.8, "Verifying Analytical Conclusions", [
        "Claim: Asterloom Systems (M03) has prepaid vs credit objection. Evidence Check: Turn 6 verbatim: 'Wait, these are prepaid? I had thought about a corporate credit card.' Verified in source turns.",
        "Claim: Willowbank (C07) agreed to review proposal. Evidence Check: Quote: 'send whatever proposal you have... call on Monday.' Verified in call text.",
        "Honest Uncertainty: For C05 (StoreHarbor), the text does NOT establish contract expiration or switching intent. Marked explicitly as 'Does Not Establish'."
    ], ACCENT_BLUE)

    # -------------------------------------------------------------
    # SLIDE 5: Trade-offs, AI Usage, Limitations & Roadmap
    # -------------------------------------------------------------
    s5 = add_base_slide(
        "Trade-Offs, AI Usage, Limitations & Next Steps",
        "Transparent engineering accounting: rejected alternatives, pragmatic boundaries, and future vision",
        5
    )
    add_card(s5, 0.8, 1.8, 5.7, 4.8, "Trade-Offs & One Rejected Decision", [
        "Deliberately Rejected Decision: Auto-overwriting manual manager routes upon receiving late revisions (e.g., M07).",
        "Why Rejected: Managers have offline context (verbal escalations, executive decisions). Erasing human input destroys user trust.",
        "Accepted Trade-Off: Route stays locked to manager's choice, but UI presents a prominent 'Revision Updated' badge and diff inspector.",
        "Pre-Generated vs Dynamic Analysis: Analytical corpus pre-generated with deterministic rules; dynamic browser client executes live filtering, search, routing overrides, and export generation."
    ], ACCENT_AMBER)

    add_card(s5, 6.8, 1.8, 5.7, 4.8, "Limitations & Production Roadmap", [
        "Current Scope: 44 records on 2026-08-21; single-user browser localStorage persistence without cloud auth.",
        "Hours Spent: ~7 hours across data audit, engine design, UI development, test automation, and documentation.",
        "Roadmap Step 1: Real-time Webhook Streaming: Integrate live Zoom/Meet bot webhooks with automatic transcription pipelines.",
        "Roadmap Step 2: Automated Telephony Diagnostics: Integrate Twilio/telephony status callbacks to distinguish failed carrier handshakes from silent dials."
    ], ACCENT_BLUE)

    output_path = os.path.join(os.path.dirname(__file__), "../slides/slides.pptx")
    prs.save(output_path)
    print(f"Successfully generated PowerPoint slides: {output_path}")

if __name__ == "__main__":
    create_presentation()
